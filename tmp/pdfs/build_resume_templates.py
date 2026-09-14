from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

OUT = Path('output/resume-templates')
OUT.mkdir(parents=True, exist_ok=True)

def para(d, text, bold=False, style=None):
    p=d.add_paragraph(style=style)
    r=p.add_run(text); r.bold=bold
    return p

def heading(d,text):
    p=para(d,text.upper(),True)
    p.paragraph_format.space_before=Pt(11)
    p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.keep_with_next=True

def bullet(d,text):
    p=para(d,text,style='List Bullet')
    p.paragraph_format.space_after=Pt(3)

def entry(d,title,detail):
    p=para(d,title,True); p.paragraph_format.keep_with_next=True
    p=para(d,detail); p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_after=Pt(4)

def education(d):
    heading(d,'Education')
    entry(d,'Bachelor of Engineering (Honours), Civil | UNSW','2025 - present | Expected graduation 2028 | WAM 78')
    para(d,'High Distinction: Engineering Mechanics')
    para(d,'Selected Distinctions: Engineering Design, Engineering Construction, Civil and Environmental Engineering Computations')

def work(d,full=False):
    heading(d,'Work Experience')
    entry(d,'Selective Tutor and Administrator','[Organisation name or Self-employed, if accurate] | Feb 2025 - present')
    bullet(d,'Developed individual learning plans for primary students preparing for selective school entry, using diagnostic assessments and progress tracking.')
    bullet(d,'Coordinated weekly group writing classes, managing tutor schedules, parent communications and payments.')
    if full: bullet(d,'Supported a selective mentoring program by helping organise personalised weekly plans around student needs and schedules.')
    entry(d,'Mathematics, Physics and Chemistry Tutor | Art of Smart','Hornsby, NSW | Feb 2025 - present')
    bullet(d,'Delivered HSC and preliminary tutoring, adapting explanations and custom problem sets to student learning needs.')
    bullet(d,'Tracked student progress and communicated feedback to parents to guide learning priorities.')

def leadership(d,full=False):
    heading(d,'Leadership')
    entry(d,'Executive Secretary | Sydney Boys High School Media Team','2022 - 2024')
    bullet(d,'Led a recruitment campaign that grew active membership from 15 to over 40.')
    bullet(d,'Coordinated sound and visual production for TEDxYouth@SBHS 2022, delegating tasks across a seven-person team.')
    if full: bullet(d,'Organised filming schedules and logistics across multiple performances of the SBHS/SGHS Matilda musical.')

for kind in ['engineering','experience']:
    d=Document(); s=d.sections[0]
    s.page_width=Inches(8.27); s.page_height=Inches(11.69)
    s.top_margin=s.bottom_margin=Inches(.6)
    s.left_margin=s.right_margin=Inches(.7)
    for name in ['Normal','Title','List Bullet']:
        st=d.styles[name]; st.font.name='Calibri'; st.font.size=Pt(10.5); st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_after=Pt(2)
        st.paragraph_format.line_spacing=1.04
    d.styles['Title'].font.size=Pt(24)
    d.styles['Title'].paragraph_format.space_after=Pt(3)
    d.styles['List Bullet'].paragraph_format.left_indent=Inches(.15)
    d.styles['List Bullet'].paragraph_format.first_line_indent=Inches(-.12)
    para(d,'Christopher Santoso',style='Title')
    para(d,'Civil Engineering Student | UNSW',True)
    para(d,'+61 478 889 668 | santosochris2006@gmail.com | [LinkedIn URL]')
    if kind=='engineering':
        education(d)
        heading(d,'Engineering Projects')
        entry(d,'[Project title] | [Course or team]','[Month year] | [Your role]')
        bullet(d,'Designed [system or structure] to meet [requirements], using [engineering method or software].')
        bullet(d,'Contributed [specific work] within a team of [number]; tested [criterion] and achieved [verified result].')
        heading(d,'Technical Skills')
        para(d,'Software and coding: [Tools you have used and what you can do with them]')
        para(d,'Engineering methods: [Calculations, analysis, testing or practical skills you can demonstrate]')
        work(d)
        leadership(d)
    else:
        heading(d,'Profile')
        para(d,'UNSW civil engineering student with experience in STEM tutoring, class administration and event coordination. Led a school media recruitment campaign that grew active membership from 15 to over 40.')
        work(d,True)
        leadership(d,True)
        education(d)
        heading(d,'Skills')
        para(d,'Tutoring and communication: HSC mathematics, physics and chemistry; individual learning plans; parent feedback.')
        para(d,'Administration and coordination: Class scheduling, payment administration, event logistics and task delegation.')
        para(d,'Software: [Add relevant tools you can confidently use]')
    d.core_properties.author='Christopher Santoso'
    d.core_properties.title='Christopher Santoso Resume'
    d.save(OUT/f'Christopher-Santoso-{kind}-template.docx')
    print(OUT/f'Christopher-Santoso-{kind}-template.docx')
